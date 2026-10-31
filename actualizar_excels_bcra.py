"""
actualizar_excels_bcra.py — Atlas Macro

Datos del BCRA que no están en su API y solo se publican en planillas de Excel:
  · Balance cambiario (mercado de cambios): pago de utilidades y dividendos, compra de
    dólares de las personas (formación de activos externos), deuda y obligaciones negociables,
    inversión directa, cuenta corriente cambiaria...  (mensual, millones de US$)
  · Tipo de cambio real multilateral (ITCRM) y bilaterales con Brasil, EE.UU., China y la
    Zona Euro (promedio mensual, índice 17-12-2015 = 100).
  · Bandas cambiarias: piso y techo día por día (el BCRA publica desde enero de 2026, con los
    días que faltan del mes siguiente ya proyectados; abril-diciembre de 2025 se reconstruye
    con la regla de ese momento: arrancó en $1.000 y $1.400 y se movía 1% por mes).
  · Irregularidad (morosidad) del crédito: total, familias y empresas, por línea
    (personales, tarjetas, hipotecarios, prendarios, adelantos, documentos...) en %.

Genera series_excel.json con el mismo formato que series_ar.json (la página las suma al
Comparador). Si una planilla falla, deja las series anteriores de esa planilla.

Uso:
    python3 actualizar_excels_bcra.py
    python3 actualizar_excels_bcra.py -v
    python3 actualizar_excels_bcra.py --local excels    # lee las planillas ya bajadas en esa carpeta
"""

import io
from datetime import date
import json
import sys
import warnings
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from curl_cffi import requests

CARPETA = Path(__file__).parent
ARCHIVO_SALIDA = CARPETA / "series_excel.json"
VER = "-v" in sys.argv
LOCAL = sys.argv[sys.argv.index("--local") + 1] if "--local" in sys.argv else None
sesion = requests.Session(impersonate="chrome")

BASE = "https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/"
PLANILLAS = {
    "cambiario": ("informes/anexo-estadistico-mercado-cambios-balance-cambiario.xlsx", "balance_cambiario.xlsx"),
    "itcrm": ("itcrm-serie.xlsx", "itcrm.xlsx"),
    "credito": ("informes/informe-bancos-anexo.xlsx", "informe_bancos.xlsx"),
    "bandas": ("serie-completa-bandas-cambiarias.xlsx", "bandas.xlsx"),
}

GB, GT, GC = "Balance cambiario", "Tipo de cambio real", "Morosidad del crédito"
F_BC, F_IT, F_CR = "BCRA (balance cambiario)", "BCRA (ITCRM)", "BCRA (informe de bancos)"
USD = "millones de US$ en el mes"


def bajar(url):
    """Bytes del Excel; si el certificado del BCRA falla, reintenta sin verificarlo (son datos públicos)."""
    try:
        r = sesion.get(url, timeout=180)
    except Exception:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            r = sesion.get(url, timeout=180, verify=False)
    r.raise_for_status()
    if not r.content.startswith(b"PK"):
        raise RuntimeError("no vino un Excel (¿cambió el link?)")
    return r.content


def libro(clave):
    url, nombre = PLANILLAS[clave]
    datos = (Path(LOCAL) / nombre).read_bytes() if LOCAL else bajar(BASE + url)
    if VER:
        print(f"        {nombre}: {len(datos) / 1e6:,.1f} MB")
    return openpyxl.load_workbook(io.BytesIO(datos), read_only=True, data_only=True)


def mes(x):
    return x.strftime("%Y-%m") if isinstance(x, datetime) else None


def num(x):
    return float(x) if isinstance(x, (int, float)) else None


def serie(nombre, unidad, grupo, fuente, filas, dec=2):
    return {"nombre": nombre, "unidad": unidad, "grupo": grupo, "fuente": fuente,
            "serie": [[m, round(v, dec)] for m, v in sorted(filas.items()) if v is not None]}


# ---------------------------------------------------------------- balance cambiario
# Cada columna de la hoja "Balance Cambiario" tiene un código fijo (bal001, bal002...) en la fila de títulos.
# código → (clave, nombre, signo)    signo -1: se da vuelta para que "compra de dólares" sea positivo
COLUMNAS_BC = {
    "bal001": ("bc_cc", "Cuenta corriente cambiaria", 1),
    "bal002": ("bc_bienes", "Bienes (cobros de exportaciones − pagos de importaciones)", 1),
    "bal005": ("bc_servicios", "Servicios (neto)", 1),
    "bal009": ("bc_intereses", "Intereses (neto)", 1),
    "bal016": ("bc_div_neto", "Utilidades y dividendos (neto)", 1),
    "bal018": ("bc_dividendos", "Pago de utilidades y dividendos al exterior", 1),
    "bal024": ("bc_cf", "Cuenta financiera cambiaria", 1),
    "bal025": ("bc_ied", "Inversión directa de no residentes (neto)", 1),
    "bal028": ("bc_portafolio", "Inversión de portafolio de no residentes (neto)", 1),
    "bal031": ("bc_deuda", "Préstamos financieros, títulos de deuda y ONs (neto)", 1),
    "bal032": ("bc_deuda_ing", "Ingresos por préstamos, títulos de deuda y ONs", 1),
    "bal034": ("bc_fmi", "Operaciones con el FMI (neto)", 1),
    "bal037": ("bc_ooii", "Préstamos de organismos internacionales (neto)", 1),
    "bal040": ("bc_fae", "Compra neta de dólares del sector privado (formación de activos externos)", -1),
    "bal050": ("bc_var_res", "Variación de reservas por operaciones de cambio", 1),
}


def parte_cambiario():
    print("· Balance cambiario")
    wb = libro("cambiario")
    filas = list(wb["Balance Cambiario"].iter_rows(values_only=True))
    codigos = next(f for f in filas if f and f[0] == "bal000")
    col = {c: j for j, c in enumerate(codigos) if c}
    falta = [c for c in list(COLUMNAS_BC) + ["bal003", "bal004", "bal017", "bal023"] if c not in col]
    if falta:
        raise RuntimeError(f"no encontré las columnas {falta}")
    out = {}
    datos = [f for f in filas if mes(f[0])]
    for cod, (clave, nombre, signo) in COLUMNAS_BC.items():
        out[clave] = serie(nombre, USD, GB, F_BC, {mes(f[0]): num(f[col[cod]]) * signo for f in datos if num(f[col[cod]]) is not None}, 1)
    # controles: total = ingresos − egresos, y cuenta corriente + capital + financiera = variación de reservas
    # (si no cierran, el BCRA cambió las columnas de lugar)
    u = [num(x) or 0 for x in datos[-1]]
    for nombre, d in (("bienes", u[col["bal002"]] - (u[col["bal003"]] - u[col["bal004"]])),
                      ("utilidades", u[col["bal016"]] - (u[col["bal017"]] - u[col["bal018"]])),
                      ("total", u[col["bal001"]] + u[col["bal023"]] + u[col["bal024"]] - u[col["bal050"]])):
        if abs(d) > 1:
            raise RuntimeError(f"las columnas de {nombre} no cierran ({d:,.1f}): revisar la planilla")

    # Personas humanas: de la hoja larga con todas las operaciones por sector y concepto (en dólares)
    ph = defaultdict(float)
    for f in wb["Datos Mercado de Cambios"].iter_rows(min_row=2, values_only=True):
        if f[2] == "Personas Humanas" and str(f[5] or "").startswith("03- Compra-venta de billetes") and mes(f[1]) and num(f[3]) is not None:
            ph[mes(f[1])] -= num(f[3]) / 1e6
    if ph:
        out["bc_fae_ph"] = serie("Compra neta de dólares de las personas", USD, GB, F_BC, ph, 1)
    for k in ("bc_dividendos", "bc_fae", "bc_fae_ph", "bc_deuda"):
        if k in out and out[k]["serie"]:
            m, v = out[k]["serie"][-1]
            anio = sum(x for mm, x in out[k]["serie"] if mm.startswith(m[:4]))
            print(f"     · {out[k]['nombre'][:60]}: {m} = {v:,.0f} · en {m[:4]}: {anio:,.0f}")
    print(f"     · {len(out)} series hasta {datos[-1][0]:%Y-%m}")
    return out


# ---------------------------------------------------------------- tipo de cambio real
ITCRM = {"ITCRM": ("itcrm", "Tipo de cambio real multilateral (ITCRM)"),
         "ITCRB Brasil": ("itcrb_brasil", "Tipo de cambio real con Brasil"),
         "ITCRB Estados Unidos": ("itcrb_eeuu", "Tipo de cambio real con EE.UU."),
         "ITCRB China": ("itcrb_china", "Tipo de cambio real con China"),
         "ITCRB Zona Euro": ("itcrb_euro", "Tipo de cambio real con la Zona Euro")}


def parte_itcrm():
    print("· Tipo de cambio real (ITCRM)")
    wb = libro("itcrm")
    filas = list(wb["ITCRM y bilaterales prom. mens."].iter_rows(values_only=True))
    tit = next(f for f in filas if f and str(f[0]).strip().lower().startswith("per"))
    col = {str(c).strip(): j for j, c in enumerate(tit) if c}
    out = {}
    for t, (clave, nombre) in ITCRM.items():
        if t not in col:
            print(f"     ⚠️  no encontré la columna «{t}»")
            continue
        out[clave] = serie(nombre, "índice 17-dic-2015 = 100 (más alto = peso más barato)", GT, F_IT,
                           {mes(f[0]): num(f[col[t]]) for f in filas if mes(f[0]) and num(f[col[t]]) is not None})
    s = out["itcrm"]["serie"]
    print(f"     · ITCRM: {len(s)} meses · último {s[-1]}")
    return out


# ---------------------------------------------------------------- irregularidad del crédito
BLOQUES = {"1. Total Sector Privado": ("irr", "total"), "2. Familias": ("irr_fam", "familias"), "3. Empresas": ("irr_emp", "empresas")}
LINEAS = {"Cartera irregular total": ("", "total"), "Adelantos": ("_adelantos", "adelantos"), "Documentos": ("_documentos", "documentos"),
          "Con garantía hipotecaria": ("_hipotecarios", "hipotecarios"), "Con garantía prendaria": ("_prendarios", "prendarios"),
          "Personales": ("_personales", "personales"), "Tarjetas de crédito": ("_tarjetas", "tarjetas de crédito"),
          "Prefinanciación y financiación de exportaciones": ("_export", "prefinanciación de exportaciones"), "Otros": ("_otros", "otras líneas")}


def parte_credito():
    print("· Irregularidad del crédito")
    wb = libro("credito")
    filas = list(wb["Calidad de Cartera (por líneas)"].iter_rows(values_only=True))
    out, i = {}, 0
    while i < len(filas):
        t = str(filas[i][0] or "").strip()
        b = next((v for k, v in BLOQUES.items() if t.startswith(k)), None)
        if not b:
            i += 1
            continue
        # la fila de fechas ("En porcentaje", 2010-01, 2010-02...) y debajo una fila por línea hasta "Fuente"
        j = next(k for k in range(i, i + 6) if str(filas[k][0] or "").strip().lower().startswith("en porcentaje"))
        meses = [mes(x) for x in filas[j]]
        k = j + 1
        while k < len(filas) and filas[k][0] and not str(filas[k][0]).startswith("Fuente"):
            linea = next((v for n, v in LINEAS.items() if str(filas[k][0]).strip().startswith(n[:25])), None)
            if linea:
                clave = b[0] + linea[0]
                quien = "todo el crédito privado" if b[1] == "total" else b[1]
                nombre = f"Morosidad: {quien}" if not linea[0] else f"Morosidad de {quien}: {linea[1]}"
                out[clave] = serie(nombre, "% de la cartera en situación irregular", GC, F_CR,
                                   {m: num(v) for m, v in zip(meses, filas[k]) if m and num(v) is not None})
            k += 1
        i = k
    for c in ("irr", "irr_fam", "irr_emp"):
        if c in out:
            print(f"     · {out[c]['nombre']}: {out[c]['serie'][-1]}")
    if "irr" not in out:
        raise RuntimeError("no encontré el bloque del total")
    print(f"     · {len(out)} series")
    return out


# ---------------------------------------------------------------- bandas cambiarias
INICIO_BANDAS = date(2025, 4, 14)          # primer día hábil del régimen: piso $1.000, techo $1.400


def parte_bandas():
    print("· Bandas cambiarias")
    wb = libro("bandas")
    filas = []
    for f in wb[wb.sheetnames[0]].iter_rows(values_only=True):
        f = [x for x in (f or []) if x is not None]           # la tabla no arranca en la columna A
        if len(f) >= 3 and isinstance(f[0], datetime) and num(f[1]) and num(f[2]):
            filas.append(f)
    if not filas:
        raise RuntimeError("no encontré las filas con fechas")
    oficial = [[f[0].date().isoformat(), round(num(f[1]), 2), round(num(f[2]), 2)] for f in filas]
    # 2025: interpolación geométrica día por día entre el inicio (1.000 / 1.400) y el primer dato oficial de 2026,
    # que equivale a moverse a una tasa diaria constante (−1% y +1% por mes)
    d1 = date.fromisoformat(oficial[0][0])
    total = (d1 - INICIO_BANDAS).days
    rec, d = [], INICIO_BANDAS
    while d < d1:
        if d.weekday() < 5:
            t = (d - INICIO_BANDAS).days / total
            rec.append([d.isoformat(), round(1000 * (oficial[0][1] / 1000) ** t, 2), round(1400 * (oficial[0][2] / 1400) ** t, 2)])
        d = date.fromordinal(d.toordinal() + 1)
    hoy = date.today().isoformat()
    actual = [f for f in oficial if f[0] <= hoy][-1] if any(f[0] <= hoy for f in oficial) else oficial[0]
    print(f"     · oficial desde {oficial[0][0]} hasta {oficial[-1][0]} · hoy ({actual[0]}): piso ${actual[1]:,.2f} · techo ${actual[2]:,.2f}")
    return {"serie": rec + oficial, "oficial_desde": oficial[0][0]}


def main():
    print("Atlas Macro · planillas del BCRA (balance cambiario, ITCRM, morosidad)")
    anterior = json.load(open(ARCHIVO_SALIDA, encoding="utf-8")) if ARCHIVO_SALIDA.exists() else {"series": {}}
    salida = {"generado": datetime.now(timezone.utc).isoformat(timespec="seconds"),
              "fuente": "BCRA (planillas de Excel: balance cambiario, ITCRM e informe de bancos)", "series": {}}
    ok = 0
    for prefijos, funcion in ((("bc_",), parte_cambiario), (("itcr",), parte_itcrm), (("irr",), parte_credito)):
        try:
            salida["series"].update(funcion())
            ok += 1
        except Exception as e:
            print(f"     ❌ {str(e)[:150]} (dejo el dato anterior)")
            salida["series"].update({k: v for k, v in anterior["series"].items() if k.startswith(prefijos)})
    try:
        salida["bandas"] = parte_bandas()
        ok += 1
    except Exception as e:
        print(f"     ❌ bandas: {str(e)[:150]} (dejo el dato anterior)")
        if "bandas" in anterior:
            salida["bandas"] = anterior["bandas"]
    if ok == 0:
        sys.exit("❌ No se pudo leer ninguna planilla: dejo el series_excel.json anterior")
    with open(ARCHIVO_SALIDA, "w", encoding="utf-8") as archivo:
        json.dump(salida, archivo, ensure_ascii=False, separators=(",", ":"))
    print(f"✅ series_excel.json listo · {len(salida['series'])} series · {ARCHIVO_SALIDA.stat().st_size / 1024:,.0f} KB")


if __name__ == "__main__":
    main()
